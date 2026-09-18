# -*- coding: utf-8 -*-
"""합성 테스트 픽스처 생성기 — 실제 양식 구조를 축소 재현한 가상 회사 자료.

사용: python tests/make_fixtures.py <대상폴더>
생성물: 전기 일반조서/계정별조서/정산표/DSD + 당기 전산자료(Raw)
(실데이터 아님 — '주식회사 테스트상사', 전기 제21기/2025)
"""
import os
import sys
import zipfile
from datetime import datetime

from openpyxl import Workbook

PRIOR_FYE = datetime(2025, 12, 31)


def make_general_wp(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Index'
    ws['A1'] = '주식회사 테스트상사'
    ws['A2'] = '감사조서 Index'
    ws['B3'] = PRIOR_FYE
    ws['A5'], ws['B5'] = '조서번호', '감사조서'
    ws['A6'], ws['B6'], ws['D6'] = '1100', '계약전 위험평가', datetime(2025, 2, 14)
    for name in ('1100', '2700', '8600'):
        s = wb.create_sheet(name)
        s['A2'] = f'{name} 조서'
        s['A4'], s['B4'] = '회사명', '주식회사 테스트상사'
        s['C4'], s['D4'], s['E4'], s['F4'] = '작성자', '홍길동', '일자', datetime(2025, 8, 13)
    wb.save(path)


def make_account_wp(path):
    wb = Workbook()
    ws = wb.active
    ws.title = '입증감사절차'
    ws['A1'] = '입 증 감 사 절 차'
    ws['A5'], ws['B5'], ws['C5'], ws['D5'] = '인덱스', '계정과목', '담당자', '종료일'
    rows = [('A', '현금및현금성자산'), ('C', '매출채권'), ('P', '매출')]
    for i, (idx, acct) in enumerate(rows, start=6):
        ws.cell(row=i, column=1, value=idx)
        ws.cell(row=i, column=2, value=acct)
        ws.cell(row=i, column=3, value='홍길동')
        ws.cell(row=i, column=4, value=datetime(2017, 3, 25))
    for s in ('A', 'C', 'P'):
        sh = wb.create_sheet(s)
        sh['A1'] = '입 증 감 사 절 차'
        sh['B3'] = PRIOR_FYE
    wb.save(path)


def make_trial_sheet(path):
    wb = Workbook()
    ar = wb.active
    ar.title = 'AR'
    ar['K1'] = '정산표'
    ar['L2'] = '주식회사 테스트상사'
    ar['L3'] = PRIOR_FYE
    ar['E4'] = '과 목'
    ar['F4'] = '제 21 (당)기  [2025/01/01 ~ 2025/12/31]'
    ar['H4'] = '제 20 (전)기  [2024/01/01 ~ 2024/12/31]'
    ar['F5'], ar['H5'] = '금 액', '금 액'
    ar['K5'], ar['L5'] = '공시용계정과목', '회사제시계정과목'
    ar['M5'], ar['N5'], ar['O5'], ar['P5'], ar['Q5'] = '2024 감사후', 2025, '차변', '대변', '2025 감사후'
    # (계정명, F세부, G본란, H전기세부, I전기본란, M, N수식, Q값)
    data = [
        ('현금',       None, 90,  None, 80,  80,  '=F6+G6',  90),
        ('외상매출금', 110,  None, 95,  None, 95,  '=F7+G7',  110),
        ('대손충당금', 10,   100,  5,   90,  -5,  '=-F8',    -10),
        ('재고자산',   None, 95,  None, 70,  70,  '=F9+G9',  95),
        ('외상매입금', None, 100, None, 90,  90,  '=F10+G10', 100),
        ('자본금',     None, 175, None, 150, 150, '=F11+G11', 175),
        ('매출',       None, 480, None, 420, 420, '=F12+G12', 480),
        ('매출원가',   None, 290, None, 260, 260, '=F13+G13', 290),
    ]
    for i, (nm, f, g, h, iv, m, nfml, q) in enumerate(data, start=6):
        ar.cell(row=i, column=5, value=nm)
        ar.cell(row=i, column=12, value=f'=E{i}')      # 회사제시계정과목 = 수식 (실양식 재현)
        for col, v in ((6, f), (7, g), (8, h), (9, iv), (13, m), (17, q)):
            if v is not None:
                ar.cell(row=i, column=col, value=v)
        ar.cell(row=i, column=14, value=nfml)
        ar.cell(row=i, column=15, value=3 if nm == '매출' else None)   # 전기 수정분개 잔존
    sad = wb.create_sheet('SAD')
    sad['A1'] = 'AJE Summary'
    cf = wb.create_sheet('CF')
    cf['A3'], cf['C3'], cf['D3'] = '과 목', '제 21기', '제 20기'
    cf['A4'], cf['C4'], cf['D4'] = 'Ⅰ. 영업활동현금흐름', 60, 45
    fn = wb.create_sheet('FN')
    fn['A2'], fn['B2'], fn['C2'] = '계정과목', '당  기  말', '전  기  말'
    fn['A3'], fn['B3'], fn['C3'] = '단기금융상품', 30, 20
    wb.save(path)


def make_raw(path):
    wb = Workbook()
    rq = wb.active
    rq.title = 'RQ'
    rq['A2'] = 'FY2026 테스트상사 기말감사 요청자료'
    bs = wb.create_sheet('BS')
    bs['A1'], bs['B1'], bs['D1'] = '과목', '제 22(당)기[2026/01/01 ~ 2026/12/31]', '제 21(전)기'
    rows = [('Ⅰ.유동자산', None, 300), ('현금', None, 100), ('외상매출금', 120, None),
            ('대손충당금', 20, 100), ('재고자산', None, 100), ('자산총계', None, 300),
            ('외상매입금', None, 120), ('부채총계', None, 120),
            ('자본금', None, 180), ('자본총계', None, 180), ('부채와자본총계', None, 300)]
    for i, (nm, b, c) in enumerate(rows, start=3):
        bs.cell(row=i, column=1, value=nm)
        if b is not None:
            bs.cell(row=i, column=2, value=b)
        if c is not None:
            bs.cell(row=i, column=3, value=c)
    pl = wb.create_sheet('PL')
    pl['A1'], pl['B1'] = '과목', '제 22(당)기'
    for i, (nm, c) in enumerate([('매출', 500), ('매출원가', 320), ('당기순이익', 60)], start=3):
        pl.cell(row=i, column=1, value=nm)
        pl.cell(row=i, column=3, value=c)
    wb.save(path)


def make_dsd(path):
    contents = '''<?xml version="1.0" encoding="utf-8"?>
<DOCUMENT xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<DOCUMENT-HEADER AEXT-CLASS="Y">
<DOCUMENT-NAME ACODE="00760">감사보고서</DOCUMENT-NAME>
<COMPANY-NAME AREGCIK="00000000">주식회사 테스트상사</COMPANY-NAME>
</DOCUMENT-HEADER>
<BODY>
<TABLE><TBODY>
<TR><TD>제 21(당) 기</TD><TD>제 20(전) 기</TD></TR>
<TR><TD>제 21 기 2025년 12월 31일 현재</TD></TR>
<TR><TE ACODE="11200000040000" ADELIM="0">Ⅰ.유동자산</TE>
<TE ACODE="11200000040000" ADELIM="1">100</TE>
<TE ACODE="11200000040000" ADELIM="2">285</TE>
<TE ACODE="11200000040000" ADELIM="3">90</TE>
<TE ACODE="11200000040000" ADELIM="4">240</TE></TR>
<TR><TD>이 감사보고서는 감사보고서일(2026년 3월 24일) 현재로 유효한 것입니다.</TD></TR>
</TBODY></TABLE>
</BODY>
</DOCUMENT>'''
    meta = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<METAINFO>'
            '<GENERATOR schema="dart4.xsd" editver="5.107"/>'
            '<DOCUMENT-HEADER regcik="00000000" regname="주식회사 테스트상사"/></METAINFO>')
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('contents.xml', contents)
        z.writestr('meta.xml', meta)


def make(folder):
    os.makedirs(os.path.join(folder, 'rawdata'), exist_ok=True)
    make_general_wp(os.path.join(folder, 'wp_일반조서_2025.xlsx'))
    make_account_wp(os.path.join(folder, 'wp_4000_계정별조서_2025.xlsx'))
    make_trial_sheet(os.path.join(folder, 'wp_8600A_정산표_2025.xlsx'))
    make_dsd(os.path.join(folder, 'FY2025_테스트상사_DSD.dsd'))
    make_raw(os.path.join(folder, 'rawdata', '2026_전산자료.xlsx'))
    print(f'픽스처 생성 완료: {folder}')


if __name__ == '__main__':
    make(sys.argv[1])
